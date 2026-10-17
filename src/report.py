import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pathlib import Path
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

proc = Path("data/processed")
out = Path("outputs")
out.mkdir(parents=True, exist_ok=True)

summary = pd.read_csv(proc / "variance_by_store_category.csv")
weekly = pd.read_csv(proc / "budget_vs_actual_weekly.csv")

# 1. Flag rule
RED, AMBER = -5.0, -2.0

def flag(v):
    if v <= RED:
        return "RED"
    if v <= AMBER:
        return "AMBER"
    return "GREEN"

summary["flag"] = summary["revenue_var_pct"].apply(flag)
summary["label"] = summary["store_id"] + " " + summary["category"]

# 2. Chart 1: budget vs actual revenue by store
by_store = (summary.groupby(["store_id", "city"], as_index=False)
            [["budget_revenue", "actual_revenue"]].sum())
names = by_store["store_id"] + "\n" + by_store["city"]
x = list(range(len(by_store)))
w = 0.4
fig, ax = plt.subplots(figsize=(12, 6))
ax.bar([i - w / 2 for i in x], by_store["budget_revenue"] / 1e6, w, label="Budget")
ax.bar([i + w / 2 for i in x], by_store["actual_revenue"] / 1e6, w, label="Actual")
ax.set_xticks(x)
ax.set_xticklabels(names)
ax.set_ylabel("Revenue (millions)")
ax.set_title("FY2026 Budget vs Actual Revenue by Store")
ax.legend()
fig.tight_layout()
fig.savefig(out / "chart_1_store_budget_vs_actual.png", dpi=150)
plt.close(fig)

# 3. Chart 2: variance % by store and category, coloured by flag
s = summary.sort_values("revenue_var_pct", ascending=False)
colors = s["flag"].map({"RED": "#d9534f", "AMBER": "#f0ad4e", "GREEN": "#5cb85c"})
fig, ax = plt.subplots(figsize=(10, 8))
ax.barh(s["label"], s["revenue_var_pct"], color=colors)
ax.axvline(RED, color="red", linestyle="--", linewidth=1)
ax.axvline(AMBER, color="orange", linestyle="--", linewidth=1)
ax.set_xlabel("Revenue variance vs budget (%)")
ax.set_title("Variance by Store and Category (red line = -5%, orange line = -2%)")
fig.tight_layout()
fig.savefig(out / "chart_2_variance_pct.png", dpi=150)
plt.close(fig)

# 4. Chart 3: weekly trend, total budget vs actual
trend = weekly.groupby("week_of_year")[["budget_revenue", "actual_revenue"]].sum() / 1e6
fig, ax = plt.subplots(figsize=(12, 5))
ax.plot(trend.index, trend["budget_revenue"], label="Budget")
ax.plot(trend.index, trend["actual_revenue"], label="Actual")
ax.set_xlabel("Week of fiscal year")
ax.set_ylabel("Revenue (millions)")
ax.set_title("Weekly Budget vs Actual (All Stores)")
ax.legend()
fig.tight_layout()
fig.savefig(out / "chart_3_weekly_trend.png", dpi=150)
plt.close(fig)

# 5. Colour-coded Excel report
cols = ["store_id", "city", "category", "budget_revenue", "actual_revenue",
        "revenue_var", "revenue_var_pct", "budget_margin_pct",
        "actual_margin_pct", "flag"]
fills = {
    "RED": PatternFill("solid", fgColor="F8CBAD"),
    "AMBER": PatternFill("solid", fgColor="FFE699"),
    "GREEN": PatternFill("solid", fgColor="C6E0B4"),
}
with pd.ExcelWriter(out / "variance_report.xlsx", engine="openpyxl") as writer:
    summary.sort_values("revenue_var_pct")[cols].to_excel(
        writer, sheet_name="Exceptions", index=False)
    ws = writer.sheets["Exceptions"]
    flag_col = cols.index("flag") + 1
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(row=r, column=flag_col)
        cell.fill = fills[cell.value]
    for i in range(1, len(cols) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 18
    ws.freeze_panes = "A2"

# 6. Print a quick check
print("Flag counts:")
print(summary["flag"].value_counts().to_string())
print()
print("Files saved in outputs:")
print("  chart_1_store_budget_vs_actual.png")
print("  chart_2_variance_pct.png")
print("  chart_3_weekly_trend.png")
print("  variance_report.xlsx")